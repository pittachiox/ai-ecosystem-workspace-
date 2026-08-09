import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_report_docx(filename="WTN-A06_Report.docx"):
    doc = docx.Document()
    
    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Cordia New'
    normal_style.font.size = Pt(14)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Document Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("รายงานสถาปัตยกรรมระบบและการออกแบบ API List (WTN-A06)")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    # Subtitle Metadata
    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("รายวิชา: AI Ecosystem  |  สัดส่วนคะแนน: 5%\nGitHub Repository: https://github.com/pittachiox/ai-ecosystem-workspace-.git")
    sub_run.font.size = Pt(12)
    sub_run.font.italic = True
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section Helper
    def add_heading_1(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        return p

    # Section 1
    add_heading_1("1. บทนำและวัตถุประสงค์โปรเจกต์")
    doc.add_paragraph(
        "การดำเนินงานในโปรเจกต์ WTN-A06 มีเป้าหมายหลักในการเปลี่ยนสภาพแวดล้อมการทำงานเดิม (Sandbox) "
        "ให้กลายเป็น Production-Ready AI Ecosystem Backend ที่วางโครงสร้างเป็นระเบียบ มีการแยกเลเยอร์หน้าที่การทำงานอย่างชัดเจน "
        "ตามหลัก Modular Clean Architecture เพื่อรองรับการขยายตัวของระบบในอนาคต โดยสถาปัตยกรรมนี้รองรับแอปพลิเคชัน 2 สายงานหลัก ได้แก่:"
    )
    p_ts = doc.add_paragraph(style='List Bullet')
    r = p_ts.add_run("Time-Series Application: ")
    r.bold = True
    p_ts.add_run("ระบบประมวลผลข้อมูลอนุกรมเวลา รองรับการทำ Resampling, Queue-based Model Training และ Forecasting")

    p_nts = doc.add_paragraph(style='List Bullet')
    r = p_nts.add_run("Non-Time Series Application: ")
    r.bold = True
    p_nts.add_run("ระบบประมวลผลข้อมูลรูปภาพ รองรับ Image Preprocessing, Object Storage (MinIO), Annotation Sync (Label Studio) และ Prediction Workflows")

    # Section 2
    add_heading_1("2. โครงสร้าง Repository จริงและการจัดหมวดหมู่")
    doc.add_paragraph("จากการตรวจสอบซอร์สโค้ดจริงใน Repository ได้ทำการย้ายงานทดลอง Sandbox เก่าเข้าสู่ legacy_labs/ และจัดโครงสร้างไฟล์หลักดังนี้:")

    code_p = doc.add_paragraph()
    code_run = code_p.add_run(
"""
.
├── README.md                      # Master Documentation นอกสุด
├── api_list_snapshot.csv          # Snapshot ตาราง API แบบ CSV
├── api_list_snapshot.xlsx         # Snapshot ตาราง API แบบ Excel
├── main.py                        # FastAPI Server Entrypoint นอกสุด
├── docker-compose.yml             # Infrastructure Services Stack
├── requirements.txt               # Dependencies Manifest
├── app/                           # Core Application Layer
│   ├── README.md
│   ├── main.py                    # App Module Entrypoint
│   ├── api/v1/                    # API Controller Routers
│   ├── core/                      # Environment Settings (config.py)
│   ├── services/                  # External Wrappers (MinIO, Label Studio, Redis)
│   ├── workers/                   # ARQ Background Job Processors
│   └── schemas/                   # Pydantic DTO Schemas
├── scripts/                       # Automation Utilities (export_openapi_to_csv.py)
├── storage/                       # Local Artifacts, Datasets, and Logs
├── diagrams/                      # System Architecture Diagrams
├── frontend/                      # Frontend UI Components
└── legacy_labs/                   # Archived Sandbox Experiments
"""
    )
    code_run.font.name = 'Consolas'
    code_run.font.size = Pt(10)

    doc.add_paragraph("หมายเหตุ: ทุกโฟลเดอร์หลักและโฟลเดอร์ย่อยในระบบมีไฟล์ README.md กำกับหน้าที่การทำงาน (Developer Responsibilities) ครบถ้วนตามข้อกำหนด 100%")

    # Section 3
    add_heading_1("3. เทคโนโลยีและ Infrastructure Stack")
    add_heading_2("3.1 Infrastructure Services (docker-compose.yml)")

    # Infrastructure Table
    infra_table = doc.add_table(rows=5, cols=4)
    infra_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    infra_headers = ["Service", "Container Image", "Port Mapping", "Role & Functionality"]
    infra_data = [
        ["redis", "redis:8.8.0-alpine", "6380:6379", "Message Broker และ Task Queue Backend สำหรับ ARQ Workers"],
        ["postgres_db", "postgres:15-alpine", "5432:5432", "Database Backend สำหรับ Label Studio"],
        ["label-studio-app", "heartexlabs/label-studio:latest", "8080:8080", "Data Annotation Platform"],
        ["minio", "minio/minio:latest", "9000 / 9001", "Object Storage จัดเก็บ Datasets และ Model Artifacts"]
    ]

    hdr_cells = infra_table.rows[0].cells
    for i, title in enumerate(infra_headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True

    for r_idx, row_data in enumerate(infra_data):
        row_cells = infra_table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = val

    add_heading_2("3.2 Python Libraries (requirements.txt)")
    doc.add_paragraph("ไลบรารีหลักที่ใช้งานในระบบประกอบด้วย: fastapi, uvicorn, minio, label-studio-sdk, redis, arq, pandas, openpyxl, pydantic, pydantic-settings, และ requests")

    # Section 4
    add_heading_1("4. Asynchronous Background Workers (ARQ Layer)")
    doc.add_paragraph("เพื่อป้องกันปัญหา HTTP Timeout เมื่อต้องทำ ML Training หรือ Data Processing งานหนัก ระบบได้ออกแบบ ARQ Worker Functions ควบคู่กับ Redis Queue ดังนี้:")
    
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run("Time-Series Tasks (app/workers/timeseries_worker.py): ")
    r.bold = True
    p.add_run("ประกอบด้วย resample_timeseries (ปรับความถี่ข้อมูล), train_timeseries_model (จำลองการ Train โมเดล) และ forecast_timeseries (คำนวณผลการคาดการณ์อนาคต)")

    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run("Non-Time Series Tasks (app/workers/nontimeseries_worker.py): ")
    r.bold = True
    p.add_run("ประกอบด้วย preprocess_images (จัดการ Image Augmentation), train_nontimeseries_model (จำลองการ Train โมเดล) และ predict_image_class (คำนวณผล Inference)")

    # Section 5
    add_heading_1("5. รายการ Backend APIs ที่เปิดใช้งานจริง (Live API List)")

    api_table = doc.add_table(rows=14, cols=5)
    api_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    api_headers = ["Category", "Method", "Path", "Parameters", "Description"]
    api_data = [
        ["System", "GET", "/", "none", "Root Landing Endpoint"],
        ["System", "GET", "/health", "none", "Service Health Check"],
        ["Storage", "POST", "/storage/upload", "bucket_name, object_name, file_path", "อัปโหลดไฟล์เข้า MinIO Bucket"],
        ["Storage", "GET", "/storage/files", "bucket_name (opt)", "ดึงรายชื่อ Objects ใน Bucket"],
        ["Storage", "GET", "/storage/presigned-url", "bucket_name, object_name", "สร้าง Presigned Temporary URL"],
        ["Labeling", "POST", "/ls/projects", "title, description, label_config", "สร้าง Project ใหม่ใน Label Studio"],
        ["Labeling", "POST", "/ls/import-data", "project_id, image_url, metadata", "นำเข้า Task รูปภาพเข้า Label Studio"],
        ["Labeling", "GET", "/ls/export-labels", "project_id, export_type", "ส่งออก Label Annotations"],
        ["Time-Series", "POST", "/timeseries/process", "data, sample_rate, window_size", "ส่งคิว Job ทำ Preprocessing"],
        ["Time-Series", "POST", "/timeseries/train", "model_type, dataset_size, horizon", "ส่งคิว Job รัน Training โมเดล"],
        ["Time-Series", "POST", "/timeseries/forecast", "payload object", "ส่งคิว Job คำนวณ Forecast Inference"],
        ["Non-Time-Series", "POST", "/nontimeseries/train", "model_type, dataset_size, horizon", "ส่งคิว Job รัน Training โมเดลรูปภาพ"],
        ["Non-Time-Series", "POST", "/nontimeseries/predict", "image_path, model_name", "ส่งคิว Job จำแนกประเภทรูปภาพ"]
    ]

    hdr_cells = api_table.rows[0].cells
    for i, title in enumerate(api_headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True

    for r_idx, row_data in enumerate(api_data):
        row_cells = api_table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = val

    # Section 6
    add_heading_1("6. ระบบสกัด OpenAPI Spec เป็น CSV / Excel")
    doc.add_paragraph("สคริปต์ scripts/export_openapi_to_csv.py จะดึง /openapi.json จาก FastAPI Server มาแปลงโครงสร้างแบบอัตโนมัติ และสร้างไฟล์ Snapshot ออกมา 2 รูปแบบที่ Root Directory ได้แก่ api_list_snapshot.csv (UTF-8-SIG) และ api_list_snapshot.xlsx")

    # Section 7
    add_heading_1("7. ภาพประกอบผลการดำเนินงานที่ต้องแนบเพิ่มเติม")
    img_list = [
        "1. ภาพหน้าจอ Docker Status (docker-compose ps) แสดงสถานะ Up ของทุก Container",
        "2. ภาพหน้าจอ Health Check Response (http://localhost:8000/health)",
        "3. ภาพหน้าจอ Interactive Swagger UI (http://localhost:8000/docs) แสดงครบทั้ง 4 Tags",
        "4. ภาพหน้าจอ ReDoc Documentation (http://localhost:8000/redoc)",
        "5. ภาพหน้าจอไฟล์ Snapshot api_list_snapshot.csv และ .xlsx ที่ถูกเจนออกมาจริง"
    ]
    for img_item in img_list:
        doc.add_paragraph(img_item, style='List Bullet')

    doc.save(filename)
    print(f"✅ สร้างไฟล์เอกสารสำเร็จ: {filename}")

if __name__ == "__main__":
    create_report_docx()
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_full_summary_docx(filename="AI_Ecosystem_Summary_Doc.docx"):
    doc = docx.Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Cordia New'
    normal_style.font.size = Pt(14)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    def add_title(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        p.paragraph_format.space_after = Pt(4)

    def add_subtitle(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(12)
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        p.paragraph_format.space_after = Pt(14)

    def add_h1(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)

    def add_h2(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        r_bold = p.add_run(bold_prefix)
        r_bold.bold = True
        p.add_run(text)
        p.paragraph_format.space_after = Pt(3)

    def style_table_header(row, headers):
        for i, title in enumerate(headers):
            cell = row.cells[i]
            cell.text = title
            r = cell.paragraphs[0].runs[0]
            r.font.bold = True
            r.font.size = Pt(11)

    # Document Header
    add_title("คู่มือสรุปสถาปัตยกรรมระบบ และเตรียมสอบ Midterm: AI Ecosystem")
    add_subtitle("สรุปครบถ้วน: File Organization, Docker, MinIO, Label Studio, PostgreSQL, Redis, REST API และ ARQ Workers")

    # Section 1
    add_h1("1. การจัดโครงสร้างไฟล์และสถาปัตยกรรม (File Organization & Clean Architecture)")
    doc.add_paragraph(
        "การจัดโครงสร้างไฟล์ในโปรเจกต์ WTN-A06 ใช้หลักการ Modular Clean Architecture โดยย้ายโค้ดทดลอง Sandbox เก่า "
        "เก็บไว้ในโฟลเดอร์ legacy_labs/ และจัดแบ่งโฟลเดอร์หลักภายใต้ app/ ตามหน้าที่การทำงาน เพื่อความง่ายในการบำรุงรักษาและขยายระบบ"
    )
    
    code_p = doc.add_paragraph()
    code_run = code_p.add_run(
"""
.
├── README.md                      # Master Documentation สรุปภาพรวมโปรเจกต์
├── api_list_snapshot.csv          # Snapshot ตาราง API รายการสด (CSV)
├── api_list_snapshot.xlsx         # Snapshot ตาราง API รายการสด (Excel)
├── main.py                        # Entrypoint สตาร์ท FastAPI Server หลัก
├── docker-compose.yml             # Infrastructure Stack Config
├── requirements.txt               # Dependencies Manifest
├── app/                           # Application Core Layer
│   ├── README.md                  # เอกสารกำกับหน้าที่ App Layer
│   ├── main.py                    # App Module Initialization
│   ├── api/v1/                    # API Controller Routers (storage, ls, timeseries, nontimeseries)
│   ├── core/                      # Configuration Management (config.py)
│   ├── services/                  # External Service Wrappers (minio, label_studio, redis)
│   ├── workers/                   # ARQ Background Job Processors (timeseries, nontimeseries)
│   └── schemas/                   # Pydantic DTO Schemas
├── scripts/                       # Automation Scripts (export_openapi_to_csv.py)
├── storage/                       # Local Storage (artifacts/, data/, log/)
├── diagrams/                      # System Architecture Diagrams
└── legacy_labs/                   # Sandbox Experiments ที่ถูก Archive
"""
    )
    code_run.font.name = 'Consolas'
    code_run.font.size = Pt(9.5)

    add_h2("เหตุผลการจัดโครงสร้างไฟล์แบบนี้ (สำหรับเขียนตอบข้อสอบ):")
    add_bullet("Separation of Concerns (การแยกหน้าที่): ", "แยก Controller (api/), Business Logic (services/), Background Tasks (workers/) และ Data Models (schemas/) ออกจากกันชัดเจน")
    add_bullet("Maintainability & Scalability: ", "เมื่อระบบขยายใหญ่ขึ้น สามารถเพิ่ม Router หรือ Worker ใหม่ได้โดยไม่กระทบโค้ดเดิม")
    add_bullet("Documentation Coverage: ", "มีไฟล์ README.md กำกับในทุกโฟลเดอร์เพื่อบอก Developer Responsibilities ตามมาตรฐานสถาปัตยกรรมระดับ Production")

    # Section 2
    add_h1("2. สรุปความรู้ 6 เทคโนโลยีหลักใน AI Ecosystem")

    # Docker
    add_h2("2.1 Docker & Docker Compose")
    add_bullet("หน้าที่หลัก: ", "จำลองสภาพแวดล้อม (Containerization) และควบคุมบริการทั้งหมดในระบบ (Orchestration)")
    add_bullet("บทบาทในระบบ: ", "ครอบและรันบริการหลังบ้าน (Redis, Postgres, Label Studio, MinIO) ให้ทำงานร่วมกันได้อย่างเป็นระเบียบ ผ่านไฟล์ docker-compose.yml")
    add_bullet("คีย์เวิร์ดสำคัญ: ", "Environment Isolation, Reproducibility (แก้ปัญหา Work on my machine), Port Mapping, Volumes")

    # MinIO
    add_h2("2.2 MinIO (Object Storage)")
    add_bullet("หน้าที่หลัก: ", "ระบบจัดเก็บไฟล์ขนาดใหญ่ที่ไม่ใช่ตาราง (Unstructured Data) รองรับโปรโตคอล S3 Compatible")
    add_bullet("บทบาทในระบบ: ", "เก็บไฟล์ภาพ/สื่อมัลติมีเดียดิบ, Dataset อนุกรมเวลา และไฟล์ผลลัพธ์โมเดล (Model Artifacts เช่น .pkl, .pt)")
    add_bullet("คีย์เวิร์ดสำคัญ: ", "High Scalability, Unstructured Data, Presigned URL (การสร้างลิงก์เข้าถึงไฟล์ชั่วคราวอย่างปลอดภัย)")

    # Label Studio
    add_h2("2.3 Label Studio (Data Annotation Platform)")
    add_bullet("หน้าที่หลัก: ", "เครื่องมือติดป้ายกำกับข้อมูล (Data Labeling) เพื่อสร้าง Ground Truth สำหรับนำไป Train โมเดล Machine Learning")
    add_bullet("บทบาทในระบบ: ", "ดึงไฟล์ภาพจาก MinIO มาให้ผู้ใช้ติด Bounding Box หรือ Class Tagging ก่อน Export ออกมาเป็น COCO/JSON format")
    add_bullet("คีย์เวิร์ดสำคัญ: ", "Human-in-the-loop, Ground Truth Generation, COCO/JSON Export")

    # PostgreSQL
    add_h2("2.4 PostgreSQL (Relational Database)")
    add_bullet("หน้าที่หลัก: ", "ฐานข้อมูลเชิงสัมพันธ์ (RDBMS) จัดเก็บข้อมูลที่มีโครงสร้างตารางชัดเจน มีความน่าเชื่อถือสูง")
    add_bullet("บทบาทในระบบ: ", "ทำหน้าที่เป็น Database Backend ให้กับ Label Studio เพื่อบันทึกข้อมูล Users, Projects และ Annotations Metadata")
    add_bullet("คีย์เวิร์ดสำคัญ: ", "ACID Compliance, Structured Data, Database Backend for Label Studio")

    # Redis
    add_h2("2.5 Redis (In-Memory Key-Value Store)")
    add_bullet("แบบง่าย: ", "เปรียบเหมือน 'กระดานโพสต์อิทบนโต๊ะ' ฝากข้อมูลเร็วเพราะอยู่ใน RAM ไม่ต้องเดินไปเปิดตู้เอกสาร (Harddisk)")
    add_bullet("เชิงเทคนิค: ", "In-Memory Key-Value Data Store ทำหน้าที่เป็น Message Broker และ Task Queue สื่อสารแบบ Microseconds")
    add_bullet("บทบาทในระบบ: ", "รับฝากคิวงานหนักๆ (เช่น Train Model) จาก FastAPI ส่งต่อให้ ARQ Worker ดึงไปทำข้างหลัง")
    add_bullet("คีย์เวิร์ดสำคัญ: ", "In-Memory Speed, Message Broker, Task Queue, Non-blocking Queue")

    # API & Workers
    add_h2("2.6 REST API (FastAPI) & Background Workers (ARQ)")
    add_bullet("FastAPI: ", "เปิด Endpoints ตามมาตรฐาน REST API/OpenAPI รองรับ HTTP Methods (GET, POST) แบ่งตาม Tag ชัดเจน")
    add_bullet("ARQ Workers: ", "ตัวประมวลผลงานหนักแบบ Asynchronous (Background Worker) ดึงงานจาก Redis Queue ไปทำ ป้องกัน API Timeout (HTTP 504)")

    # Section 3
    add_h1("3. ตารางคำสั่ง Docker & Docker Compose สำคัญ")
    
    cmd_table = doc.add_table(rows=8, cols=2)
    cmd_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["คำสั่ง (Command)", "หน้าที่การทำงาน และคำอธิบาย"]
    cmd_data = [
        ["docker-compose up -d", "สร้างและสตาร์ทคอนเทนเนอร์ทั้งหมดใน docker-compose.yml แบบ Background (-d)"],
        ["docker-compose down", "หยุดการทำงาน และลบคอนเทนเนอร์/เครือข่ายทั้งหมดที่สร้างโดย Compose"],
        ["docker-compose ps", "เช็กสถานะการทำงานของคอนเทนเนอร์ทั้งหมดใน Stack (Up / Exit)"],
        ["docker-compose logs -f", "ดึง Log การทำงานของทุกคอนเทนเนอร์มาดูสดๆ แบบ Real-time"],
        ["docker ps / docker ps -a", "ดูรายการคอนเทนเนอร์ที่กำลังรันอยู่ / ดูคอนเทนเนอร์ทั้งหมดในเครื่อง"],
        ["docker exec -it <id> bash", "แทรกตัวเข้าไปพิมพ์คำสั่งใน Terminal ข้างในคอนเทนเนอร์"],
        ["kill -9 $(lsof -t -i:8000)", "สั่งปิด Process ที่รันค้างอยู่บนพอร์ต 8000 บน macOS/Linux"]
    ]
    style_table_header(cmd_table.rows[0], headers)
    for r_idx, row in enumerate(cmd_data):
        row_cells = cmd_table.rows[r_idx + 1].cells
        row_cells[0].text = row[0]
        row_cells[1].text = row[1]

    # Section 4
    add_h1("4. ภาพรวมกระบวนการประมวลผลข้อมูล (Data Lifecycle Flow ใน AI Ecosystem)")
    doc.add_paragraph("หากข้อสอบถามถึงลำดับการทำงานของระบบตั้งแต่ต้นจนจบ ให้เขียนเรียงลำดับเป็น 5 ขั้นตอนดังนี้:")
    add_bullet("1. Data Ingestion: ", "ผู้ใช้ยิงอัปโหลดไฟล์ข้อมูลดิบผ่าน FastAPI Endpoint (/storage/upload)")
    add_bullet("2. Object Storage: ", "MinIO Service Wrapper บันทึกไฟล์ลงใน MinIO Bucket")
    add_bullet("3. Data Labeling: ", "Label Studio ดึงไฟล์ภาพจาก MinIO มาติด Label (โดยมี PostgreSQL บันทึกข้อมูลโครงการอยู่หลังบ้าน)")
    add_bullet("4. Async Model Training: ", "ผู้ใช้สั่ง Train โมเดลผ่าน FastAPI (/timeseries/train หรือ /nontimeseries/train) โดย API จะส่ง Job ไปฝากที่ Redis Queue จากนั้น ARQ Worker จะดึงไป Train ใน Background")
    add_bullet("5. Artifact Saving & Inference: ", "เมื่อ Train เสร็จ โมเดลจะถูกบันทึกกลับไปเก็บใน MinIO เป็น Model Artifact และเปิด Endpoint รับคำขอ Prediction/Forecast")

    doc.save(filename)
    print(f"✅ สร้างเอกสารฉบับสมบูรณ์สำเร็จ: {filename}")

if __name__ == "__main__":
    create_full_summary_docx()
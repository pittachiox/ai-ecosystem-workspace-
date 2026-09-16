#!/usr/bin/env python3
"""Simple Markdown to DOCX converter for assignment report.
- Keeps headings, paragraphs, bullet lists, numbered lists, code blocks.
- Replaces image references with placeholder paragraphs indicating where to insert images.

Note: This is a lightweight converter and won't render complex Markdown features.
"""
import re
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

import os

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
MD_PATH = os.path.join(BASE_DIR, "ASSIGNMENT_09_REPORT.md")
OUT_PATH = os.path.join(BASE_DIR, "ASSIGNMENT_09_REPORT.docx")

image_re = re.compile(r"!\[(.*?)\]\((.*?)\)")
link_re = re.compile(r"\[(.*?)\]\((.*?)\)")


def add_code_block(doc, code_lines):
    p = doc.add_paragraph()
    run = p.add_run("\n".join(code_lines))
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    # Set East Asia font to keep consistent on mac/Windows
    try:
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')
    except Exception:
        pass


def md_to_docx(md_path, out_path):
    doc = Document()
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code = False
    code_lines = []
    for raw in lines:
        line = raw.rstrip('\n')
        if line.strip().startswith('```'):
            if in_code:
                # flush code block
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        # Headings
        if line.startswith('#'):
            m = re.match(r"^(#+)\s*(.*)$", line)
            if m:
                level = len(m.group(1))
                text = m.group(2).strip()
                level = min(level, 3)
                doc.add_heading(text, level=level-1 if level>0 else 0)
            continue

        # Horizontal rule
        if line.strip() in ('---', '***', '___'):
            doc.add_paragraph('')
            continue

        # Image
        m = image_re.search(line)
        if m:
            alt = m.group(1) or 'image'
            doc.add_paragraph(f'[Image placeholder: {alt}]')
            # keep any trailing text
            trailing = image_re.sub('', line).strip()
            if trailing:
                doc.add_paragraph(trailing)
            continue

        # Lists (bullets)
        if re.match(r"^\s*[-*+]\s+", line):
            item = re.sub(r"^\s*[-*+]\s+", '', line).strip()
            p = doc.add_paragraph(item, style='List Bullet')
            continue

        # Numbered lists
        if re.match(r"^\s*\d+\.\s+", line):
            item = re.sub(r"^\s*\d+\.\s+", '', line).strip()
            p = doc.add_paragraph(item, style='List Number')
            continue

        # Links - replace with text (text (url))
        line = link_re.sub(lambda mm: f"{mm.group(1)} ({mm.group(2)})", line)

        # Normal paragraph
        if line.strip() == '':
            doc.add_paragraph('')
        else:
            doc.add_paragraph(line)

    # flush if code block left open
    if in_code and code_lines:
        add_code_block(doc, code_lines)

    doc.save(out_path)


if __name__ == '__main__':
    md_to_docx(MD_PATH, OUT_PATH)
    print(f"Wrote {OUT_PATH}")
